import os
import zipfile
from datetime import datetime
from zoneinfo import ZoneInfo

from flask import Flask, render_template, request, send_file
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# =========================
# CONFIGURACIÓN GENERAL
# =========================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
OUTPUT_DIR = os.path.join(BASE_DIR, "outputs")
PLANTILLA_WORD = os.path.join(BASE_DIR, "templates_word", "AUTOR.docx")

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

app = Flask(__name__)

FUENTE = "Arial"
TAM_FUENTE = Pt(10)

# =========================
# UTILIDADES WORD
# =========================

def aplicar_borde_doble(tabla):
    tbl = tabla._tbl
    tblPr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "double")
        borde.set(qn("w:sz"), "6")
        borde.set(qn("w:space"), "0")
        borde.set(qn("w:color"), "000000")
        borders.append(borde)

    tblPr.append(borders)


def centrar_tabla(tabla):
    tblPr = tabla._tbl.tblPr
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblPr.append(jc)


def aplicar_estilo(run, bold=False):
    run.bold = bold
    run.font.name = FUENTE
    run.font.size = TAM_FUENTE


# =========================
# PARSEO TXT
# =========================

def leer_txt(path):
    with open(path, "r", encoding="latin-1", errors="ignore") as f:
        return f.read()


def parsear_registros(texto):
    registros = []
    actual = {}

    for linea in texto.splitlines():
        linea = linea.strip()
        if not linea:
            continue

        if linea.startswith("MFN:"):
            if actual:
                registros.append(actual)
            actual = {"MFN": linea.replace("MFN:", "").strip()}
        else:
            if "\t" in linea:
                k, v = linea.split("\t", 1)
                actual[k.strip()] = v.strip()

    if actual:
        registros.append(actual)

    return registros


# =========================
# GENERACIÓN WORD
# =========================


def limpiar_documento(doc):
    body = doc._element.body
    for child in list(body):
        body.remove(child)


def agregar_ficha(doc, registro):
    filas = [
        ("SIGNATURA TOPOGRAFICA", registro.get("SIGNATURA TOPOGRAFICA", "")),
        ("AUTOR PRINCIPAL", registro.get("AUTOR PRINCIPAL", "")),
        ("TITULO / SUBTITULO", registro.get("TITULO/SUBTITULO", "")),
        ("MENCION RESPONSABILIDAD", registro.get("MENCION RESPONSABILIDAD", "")),
    ]

    if registro.get("EDICION"):
        filas.append(("EDICION", registro.get("EDICION")))

    filas.extend([
        ("IMPRENTA", registro.get("IMPRENTA", "")),
        ("DESCRIPCION FISICA", registro.get("DESCRIPCION FISICA", "")),
    ])

    tabla = doc.add_table(rows=len(filas), cols=2)
    tabla.autofit = False

    # 20cm - 4cm - 4cm = 12cm
    tabla.columns[0].width = Cm(4)
    tabla.columns[1].width = Cm(8)

    for i, (titulo, valor) in enumerate(filas):
        c1 = tabla.cell(i, 0).paragraphs[0]
        c2 = tabla.cell(i, 1).paragraphs[0]

        r1 = c1.add_run(titulo)
        aplicar_estilo(r1, bold=True)

        r2 = c2.add_run(valor)
        aplicar_estilo(r2, bold=(i < 2))

    aplicar_borde_doble(tabla)
    centrar_tabla(tabla)

    doc.add_page_break()

def crear_word(registros, ruta_salida):
    doc = Document(PLANTILLA_WORD)

    #  eliminar TODO el contenido inicial del template
    limpiar_documento(doc)

    for i, r in enumerate(registros):
        if i > 0:
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True

        agregar_ficha(doc, r)

    doc.save(ruta_salida)


# =========================
# FLASK
# =========================

@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        archivo = request.files.get("archivo")
        if not archivo:
            return "Archivo no enviado", 400

        ruta_txt = os.path.join(UPLOAD_DIR, archivo.filename)
        archivo.save(ruta_txt)

        texto = leer_txt(ruta_txt)
        registros = parsear_registros(texto)

        if not registros:
            return "No se encontraron registros", 400

        mfn_inicio = registros[0]["MFN"]
        mfn_fin = registros[-1]["MFN"]

        ahora = datetime.now(
            ZoneInfo("America/Bogota")
        ).strftime("%Y-%m-%d_%H-%M-%S")

        nombre_word = f"fichas_{ahora}_MFN_{mfn_inicio}_{mfn_fin}.docx"
        ruta_word = os.path.join(OUTPUT_DIR, nombre_word)

        crear_word(registros, ruta_word)

        nombre_zip = nombre_word.replace(".docx", ".zip")
        ruta_zip = os.path.join(OUTPUT_DIR, nombre_zip)

        with zipfile.ZipFile(ruta_zip, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(ruta_word, arcname=nombre_word)

        return send_file(ruta_zip, as_attachment=True)

    return render_template("index.html")


# =========================
# ENTRYPOINT RENDER
# =========================

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=10000)

